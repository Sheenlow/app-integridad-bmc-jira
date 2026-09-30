"""
Formateo y generacion de archivos Excel para exportaciones.
"""
import pandas as pd

from .styles import clasificar_accion, COLORES_ACCION


def _fill_accion(valor):
    """Devuelve un PatternFill segun la categoria de accion, o None si no aplica."""
    from openpyxl.styles import PatternFill
    categoria = clasificar_accion(valor)
    if categoria is None:
        return None
    bg, _ = COLORES_ACCION[categoria]
    # openpyxl espera aRGB (8 digitos): anteponemos el alfa completo 'FF'.
    return PatternFill("solid", fgColor="FF" + bg)


def formatear_excel(writer, df: pd.DataFrame, sheet_name: str, columna_color: str | None = None):
    """
    Aplica formato al Excel exportado:
    - Auto-ajuste de columnas (robusto a DataFrames vacios)
    - Header con fondo oscuro y texto blanco
    - Freeze en primera fila
    - Bordes finos
    - Color condicional en la columna especificada
    """
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
    from openpyxl.utils import get_column_letter

    ws = writer.sheets[sheet_name]

    # Auto-ajuste de columnas
    for i, col in enumerate(df.columns, 1):
        largos = df[col].astype(str).str.len()
        max_len = largos.max() if not largos.empty else 0
        max_len = max(max_len, len(str(col)))
        ws.column_dimensions[get_column_letter(i)].width = min(max_len + 3, 50)

    # Header
    header_fill = PatternFill("solid", fgColor="1e293b")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_align = Alignment(horizontal="center", vertical="center")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align

    # Freeze
    ws.freeze_panes = "A2"

    # Bordes finos
    thin = Side(style="thin", color="cbd5e1")
    border = Border(top=thin, bottom=thin, left=thin, right=thin)
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            cell.border = border
            if cell.alignment and not cell.alignment.horizontal:
                cell.alignment = Alignment(vertical="center")

    # Color condicional en columna de accion/validacion
    if columna_color and columna_color in df.columns:
        col_idx = list(df.columns).index(columna_color) + 1
        for row in range(2, ws.max_row + 1):
            cell = ws.cell(row=row, column=col_idx)
            fill = _fill_accion(cell.value)
            if fill is not None:
                cell.fill = fill


def crear_excel(df: pd.DataFrame, sheet_name: str, columna_color: str | None = None) -> bytes:
    """Genera un archivo Excel en memoria (con formato) y devuelve sus bytes."""
    from io import BytesIO
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name=sheet_name)
        formatear_excel(writer, df, sheet_name, columna_color=columna_color)
    return buffer.getvalue()
