"""
Funciones de entrada/salida: lectura robusta de archivos Excel, CSV y HTML.
Incluye parser ZIP/XML para archivos xlsx corruptos.
"""
import logging
import pandas as pd
from io import BytesIO

logger = logging.getLogger(__name__)


def detectar_columnas_duplicadas(df: pd.DataFrame, label: str, st_module=None) -> None:
    """Advierte si el DataFrame tiene nombres de columna duplicados."""
    cols = df.columns.tolist()
    duplicados = [c for c in cols if cols.count(c) > 1]
    if duplicados:
        if st_module:
            st_module.warning(
                f"**{label}** - Se detectaron columnas duplicadas: "
                f"{', '.join(sorted(set(duplicados)))}. "
                "Pandas las renombrara con sufijos (.1, .2). "
                "Verifica el archivo fuente."
            )


def listar_hojas_excel(archivo: bytes) -> list:
    """Devuelve los nombres de las hojas de un archivo Excel."""
    from openpyxl import load_workbook
    wb = load_workbook(filename=BytesIO(archivo), read_only=True)
    hojas = wb.sheetnames
    wb.close()
    return hojas


def leer_archivo_robusto(
    uploaded_file, sheet_name=0, permitir_fallback: bool = True
) -> pd.DataFrame:
    """
    Lee un archivo en cascada: Excel (openpyxl) -> CSV (auto-detect separator) -> HTML.
    Para archivos xlsx reales (permitir_fallback=False), si openpyxl falla, intenta
    extraer los datos directamente del ZIP antes de rendirse.
    Reinicia el puntero con seek(0) antes de cada intento.
    Lanza ValueError si todos los formatos fallan.
    """
    fname = getattr(uploaded_file, "name", "<stream>")
    logger.info("Leyendo archivo: %s (sheet=%s, fallback=%s)", fname, sheet_name, permitir_fallback)
    uploaded_file.seek(0)
    try:
        return pd.read_excel(uploaded_file, sheet_name=sheet_name, engine="openpyxl")
    except Exception as exc_xlsx:
        logger.warning("openpyxl no pudo leer '%s': %s", fname, exc_xlsx)
        if not permitir_fallback:
            uploaded_file.seek(0)
            try:
                logger.info("Intentando extraccion ZIP/XML...")
                return _extraer_datos_xlsx_desde_zip(uploaded_file, sheet_name)
            except Exception as exc_zip:
                raise ValueError(
                    f"No se pudo leer '{fname}' como Excel. "
                    f"openpyxl: {exc_xlsx} | extraccion ZIP/XML: {exc_zip}"
                ) from exc_zip

    for enc in ["utf-8", "latin-1", "cp1252"]:
        try:
            uploaded_file.seek(0)
            return pd.read_csv(uploaded_file, sep=None, engine="python", encoding=enc)
        except Exception as exc_csv:
            logger.debug("CSV (%s) fallo: %s", enc, exc_csv)
            continue

    try:
        uploaded_file.seek(0)
        return pd.read_html(uploaded_file)[0]
    except Exception as exc_html:
        logger.debug("HTML fallo: %s", exc_html)
        pass

    raise ValueError(
        f"No se pudo leer el archivo '{fname}' en ningun formato "
        f"(excel, csv, html). Causa raiz (openpyxl): {exc_xlsx}"
    )


def _extraer_datos_xlsx_desde_zip(
    uploaded_file, sheet_name=0
) -> pd.DataFrame:
    """
    Extrae datos de un archivo xlsx corrupto leyendo directamente los XML
    dentro del ZIP, evitando la hoja de estilos que causa el error de openpyxl.
    Soporta sheet_name como indice (int) o nombre (str).
    """
    import zipfile
    from xml.etree import ElementTree as ET

    NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    NS_R = "http://schemas.openxmlformats.org/package/2006/relationships"

    with zipfile.ZipFile(uploaded_file) as zf:

        # --- Shared strings (soporta texto simple y rich text) ---
        strings = []
        if "xl/sharedStrings.xml" in zf.namelist():
            tree = ET.parse(zf.open("xl/sharedStrings.xml"))
            root = tree.getroot()
            for si in root.findall(f"{{{NS}}}si"):
                t = si.find(f"{{{NS}}}t")
                if t is not None and t.text:
                    strings.append(t.text)
                else:
                    # Rich text: concatenar todos los <r><t> dentro de <si>
                    textos = []
                    for r_elem in si.findall(f"{{{NS}}}r"):
                        rt = r_elem.find(f"{{{NS}}}t")
                        if rt is not None and rt.text:
                            textos.append(rt.text)
                    strings.append("".join(textos))

        # --- Resolver hoja por nombre via relaciones del workbook ---
        sheet_file = None

        # Mapa rId -> target desde workbook.xml.rels
        rid_targets = {}
        rels_path = "xl/_rels/workbook.xml.rels"
        if rels_path in zf.namelist():
            rels_tree = ET.parse(zf.open(rels_path))
            for rel in rels_tree.getroot().findall(f"{{{NS_R}}}Relationship"):
                rid = rel.get("Id")
                target = rel.get("Target")
                if rid and target:
                    rid_targets[rid] = target

        # Leer workbook.xml para obtener el target de la hoja deseada
        wb_tree = ET.parse(zf.open("xl/workbook.xml"))
        wb_root = wb_tree.getroot()
        sheets_elem = wb_root.find(f"{{{NS}}}sheets")

        if isinstance(sheet_name, str):
            for s in sheets_elem.findall(f"{{{NS}}}sheet"):
                if s.get("name") == sheet_name:
                    rid = s.get(f"{{{NS_R}}}id") or s.get("r:id")
                    target = rid_targets.get(rid, f"worksheets/sheet{s.get('sheetId', '1')}.xml")
                    sheet_file = f"xl/{target}"
                    break
            if sheet_file is None:
                sheet_file = "xl/worksheets/sheet1.xml"
        else:
            idx = int(sheet_name) if sheet_name else 0
            sheet_targets = []
            for s in sheets_elem.findall(f"{{{NS}}}sheet"):
                rid = s.get(f"{{{NS_R}}}id") or s.get("r:id")
                target = rid_targets.get(rid, f"worksheets/sheet{s.get('sheetId', '1')}.xml")
                sheet_targets.append(target)
            if 0 <= idx < len(sheet_targets):
                sheet_file = f"xl/{sheet_targets[idx]}"
            else:
                sheet_file = f"xl/worksheets/sheet{idx + 1}.xml"

        if sheet_file not in zf.namelist():
            disponibles = [n for n in zf.namelist() if "sheet" in n.lower()]
            raise FileNotFoundError(
                f"Hoja '{sheet_file}' no encontrada. "
                f"Hojas disponibles en el ZIP: {disponibles or zf.namelist()[:10]}"
            )

        # --- Parsear la hoja ---
        sheet_tree = ET.parse(zf.open(sheet_file))
        sheet_root = sheet_tree.getroot()
        sheet_data = sheet_root.find(f"{{{NS}}}sheetData")

        if sheet_data is None:
            raise ValueError(
                f"La hoja '{sheet_file}' no contiene <sheetData>."
            )

        rows_data = []
        for row_elem in sheet_data.findall(f"{{{NS}}}row"):
            row = {}
            for cell in row_elem.findall(f"{{{NS}}}c"):
                ref = cell.get("r") or ""
                col_letter = "".join(c for c in ref if c.isalpha())
                value_elem = cell.find(f"{{{NS}}}v")
                formula_elem = cell.find(f"{{{NS}}}f")
                cell_type = cell.get("t", "")

                if value_elem is not None and value_elem.text is not None:
                    if cell_type == "s":
                        idx_str = int(value_elem.text)
                        row[col_letter] = strings[idx_str] if idx_str < len(strings) else ""
                    else:
                        row[col_letter] = value_elem.text
                elif cell_type == "inlineStr":
                    is_elem = cell.find(f"{{{NS}}}is")
                    if is_elem is not None:
                        t_inline = is_elem.find(f"{{{NS}}}t")
                        row[col_letter] = t_inline.text if t_inline is not None and t_inline.text else ""
                    else:
                        row[col_letter] = ""
                elif formula_elem is not None and formula_elem.text:
                    row[col_letter] = formula_elem.text
                else:
                    row[col_letter] = None

            if row:
                rows_data.append(row)

    if not rows_data:
        raise ValueError(
            "No se encontraron datos en el archivo xlsx (extraccion ZIP)."
        )

    df = pd.DataFrame(rows_data)

    if not df.empty:
        primera = df.iloc[0].tolist()
        tiene_headers = any(
            isinstance(v, str) and v.strip() for v in primera if v is not None
        )
        if tiene_headers:
            nuevos = []
            for i, val in enumerate(primera):
                if isinstance(val, str) and val.strip():
                    nuevos.append(val.strip())
                else:
                    nuevos.append(f"Col_{i + 1}")
            df.columns = nuevos
            df = df.iloc[1:].reset_index(drop=True)
        else:
            df.columns = [f"Col_{i + 1}" for i in range(len(df.columns))]

    df = df.fillna("")

    return df


def leer_archivo_subido(archivo, label: str, st_module, hoja_preferida: str | None = None) -> pd.DataFrame | None:
    """
    Lee un archivo .csv o .xlsx subido via st.file_uploader.
    - Usa leer_archivo_robusto() con fallback Excel -> CSV -> HTML.
    - Para archivos .xlsx reales, permite elegir hoja si hay mas de una.
    - Si se especifica hoja_preferida y existe, se selecciona automaticamente.
    - Detecta columnas duplicadas y emite advertencia.
    """
    if archivo is None:
        return None

    try:
        nombre = archivo.name.lower()

        if nombre.endswith((".csv", ".xlsx")):
            hoja = 0
            es_fallback = True

            if nombre.endswith(".xlsx"):
                archivo.seek(0)
                magic = archivo.read(4)
                archivo.seek(0)
                es_zip = (magic[:2] == b"PK")

                if es_zip:
                    es_fallback = False
                    try:
                        bytes_archivo = archivo.getvalue()
                        hojas = listar_hojas_excel(bytes_archivo)
                        if hoja_preferida and hoja_preferida in hojas:
                            hoja = hoja_preferida
                        elif len(hojas) == 1:
                            hoja = hojas[0]
                        else:
                            hoja = st_module.selectbox(
                                f"Hoja a leer para **{label}**",
                                options=hojas,
                                key=f"sheet_{label}",
                            )
                    except Exception:
                        hoja = 0

            with st_module.spinner(f"Leyendo {label}..."):
                df = leer_archivo_robusto(
                    archivo, sheet_name=hoja, permitir_fallback=es_fallback
                )
        else:
            st_module.error(
                f"**{label}** - Formato no soportado: '{archivo.name}'. "
                "Usa archivos .csv o .xlsx."
            )
            return None

        if df.empty:
            st_module.error(f"**{label}** - El archivo esta vacio (0 filas).")
            return None

        detectar_columnas_duplicadas(df, label, st_module)

        st_module.toast(f"✅ {label} cargado — {df.shape[0]:,} filas", icon="✅")
        return df

    except Exception as e:
        st_module.toast(f"❌ {label} — Error de lectura", icon="❌")
        st_module.error(f"**{label}** - Error al leer el archivo: {e}")
        return None
