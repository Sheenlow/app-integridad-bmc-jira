"""
Tests para la generacion/formateo de Excel.
"""
import sys
sys.path.insert(0, ".")

import pandas as pd

from src.excel_export import crear_excel
from src.constants import COL_ACCION_SUGERIDA


def test_crear_excel_genera_bytes_con_colores():
    """crear_excel debe generar un xlsx valido y colorear las acciones."""
    df = pd.DataFrame({
        COL_ACCION_SUGERIDA: [
            "OK - Sincronizado",
            "Revisar: Falta en Jira",
            "Revisar: Sobra en Jira",
            "Revisar: Estado en incongruencia con Jira",
        ],
    })
    data = crear_excel(df, "Test", columna_color=COL_ACCION_SUGERIDA)

    assert isinstance(data, bytes)
    assert len(data) > 0
    assert data[:2] == b"PK"  # magic bytes de un .xlsx (ZIP)


def test_crear_excel_legible_por_openpyxl():
    """El xlsx generado debe poder releerse sin errores."""
    from io import BytesIO
    from openpyxl import load_workbook

    df = pd.DataFrame({COL_ACCION_SUGERIDA: ["OK - Sincronizado"]})
    data = crear_excel(df, "Test", columna_color=COL_ACCION_SUGERIDA)

    wb = load_workbook(BytesIO(data))
    assert wb.sheetnames == ["Test"]


def test_crear_excel_con_df_vacio_no_rompe():
    """Un DataFrame vacio no debe fallar en el auto-ajuste de columnas."""
    df = pd.DataFrame(columns=["A", "B"])
    data = crear_excel(df, "Vacio")

    assert isinstance(data, bytes)
    assert data[:2] == b"PK"
